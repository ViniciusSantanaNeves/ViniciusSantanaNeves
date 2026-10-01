"""Definição única da estrutura da base (colunas, listas, bairros) e da planilha.

Todos os scripts importam daqui, para que a planilha gerada, o importador de
LAI e o vinculador de fotos usem exatamente as mesmas colunas e regras.
"""
from __future__ import annotations

import os
import unicodedata
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------
LINHAS_DADOS = 3000          # linhas pré-formatadas na aba Vendas
PRIMEIRA_LINHA = 2           # linha 1 = cabeçalho
ULTIMA_LINHA = PRIMEIRA_LINHA + LINHAS_DADOS - 1
PASTA_FOTOS = "fotos"        # relativa à pasta da planilha

FONTE = "Arial"

CIDADES = ["São Sebastião", "Bertioga"]
NOME_FAIXA_BAIRROS = {"São Sebastião": "BAIRROS_SAO_SEBASTIAO", "Bertioga": "BAIRROS_BERTIOGA"}

# Bertioga: Lei Complementar Municipal nº 99/2013, Anexo 1 (grafia literal).
# São Sebastião: não há lei municipal de bairros localizada; usa-se a base
# oficial de bairros do IBGE – Censo 2022 (delimitada com a prefeitura).
BAIRROS = {
    "Bertioga": [
        ("Caiubura", "01", "RA1 Sul"),
        ("São João", "02", "RA1 Sul"),
        ("Centro", "03", "RA2 Central"),
        ("Jardim Vicente de Carvalho", "04", "RA2 Central"),
        ("Albatróz", "05", "RA2 Central"),
        ("Maitinga", "06", "RA2 Central"),
        ("Rio da Praia", "07", "RA2 Central"),
        ("Buriqui Costa Nativa (Brasfanta)", "08", "RA2 Central"),
        ("Jardim Raphael", "09", "RA2 Central — conhecido popularmente como São Rafael"),
        ("Bairro Chácaras", "10", "RA3 Média"),
        ("Vista Linda", "11", "RA3 Média"),
        ("Indaiá", "12", "RA3 Média"),
        ("Riviera", "13", "RA3 Média"),
        ("São Lourenço", "14", "RA3 Média"),
        ("Guaratuba", "15", "RA4 Norte"),
        ("Costa do Sol", "16", "RA4 Norte"),
        ("Morada da Praia", "17", "RA4 Norte"),
        ("Boraceia", "18", "RA4 Norte"),
        ("T.I. (Terras Indígenas do Rio Silveiras)", "19", "RA4 Norte"),
    ],
    "São Sebastião": [
        ("Enseada", "3550704001", "Distrito São Francisco da Praia"),
        ("Jaraguá", "3550704002", "Distrito São Francisco da Praia"),
        ("Cigarras", "3550704003", "Distrito São Francisco da Praia"),
        ("São Francisco", "3550704004", "Distrito São Francisco da Praia"),
        ("Morro do Abrigo", "3550704005", "Distrito São Francisco da Praia"),
        ("Olaria", "3550704006 / 3550704036",
         "O IBGE registra dois setores 'Olaria' (distritos São Francisco da Praia e São Sebastião)"),
        ("Arrastão", "3550704007", "Distrito São Francisco da Praia"),
        ("Pontal da Cruz", "3550704008", "Distrito São Sebastião"),
        ("Praia Deserta", "3550704009", "Distrito São Sebastião"),
        ("Porto Grande", "3550704010", "Distrito São Sebastião"),
        ("Centro", "3550704011", "Distrito São Sebastião"),
        ("Topolândia", "3550704012", "Distrito São Sebastião"),
        ("Itatinga", "3550704013", "Distrito São Sebastião"),
        ("Varadouro", "3550704014", "Distrito São Sebastião"),
        ("Pitangueiras", "3550704015", "Distrito São Sebastião"),
        ("Barequeçaba", "3550704016", "Distrito São Sebastião"),
        ("Guaecá", "3550704017", "Distrito São Sebastião"),
        ("Toque-Toque Grande", "3550704018", "Distrito São Sebastião"),
        ("Calhetas", "3550704019", "Distrito São Sebastião"),
        ("Toque-Toque Pequeno", "3550704020", "Distrito São Sebastião"),
        ("Santiago", "3550704021", "Distrito Maresias"),
        ("Paúba", "3550704022", "Distrito Maresias"),
        ("Maresias", "3550704023", "Distrito Maresias"),
        ("Boiçucanga", "3550704024", "Distrito Maresias"),
        ("Camburi", "3550704025", "Distrito Maresias"),
        ("Baleia", "3550704026", "Distrito Maresias"),
        ("Barra do Sahy", "3550704027", "Distrito Maresias"),
        ("Juqueí", "3550704028", "Distrito Maresias — grafia IBGE (também escrito Juquehy)"),
        ("Barra do Una", "3550704029", "Distrito Maresias"),
        ("Boracéia", "3550704030", "Distrito Maresias"),
        ("Canto do Mar", "3550704031", "Distrito São Francisco da Praia"),
        ("Vila Amélia", "3550704034", "Distrito São Sebastião"),
        ("Reserve Du Moulin", "3550704035", "Distrito São Francisco da Praia"),
        ("Industrial", "3550704037", "Distrito São Sebastião"),
        ("Jureia", "3550704038", "Distrito Maresias"),
        ("Engenho", "3550704039", "Distrito Maresias"),
    ],
}

FONTE_BAIRROS = {
    "Bertioga": "Lei Complementar nº 99/2013 de Bertioga, Anexo 1 — "
                "https://www.bertioga.sp.gov.br/wp/wp-content/uploads/2015/06/Lei-Complementar-99.2013.pdf",
    "São Sebastião": "IBGE, Censo 2022, malha de bairros SP_bairros_CD2022 — "
                     "https://geoftp.ibge.gov.br/organizacao_do_territorio/malhas_territoriais/"
                     "malhas_de_setores_censitarios__divisoes_intramunicipais/censo_2022/bairros/shp/UF/"
                     "SP_bairros_CD2022.zip (não foi localizada lei municipal de bairros)",
}

TIPOS = ["Casa", "Casa em condomínio", "Apartamento", "Terreno", "Sala/Loja comercial", "Galpão", "Outro"]
PADROES = ["Baixo", "Normal", "Alto", "Não informado"]
ORIGEM_PADRAO = ["Cadastro municipal (ITBI/IPTU)", "Certidão/memorial/habite-se",
                 "Vistoria própria (critério NBR 12721)", "Não informado"]
NATUREZAS = ["Compra e venda"]
FONTES = ["LAI – ITBI Prefeitura", "Certidão de matrícula (CRI)"]
ORIGEM_FOTO = ["Foto própria no local", "Google Street View (link)", "Foto da certidão/laudo", "Sem foto"]
STATUS = ["Validado", "Pendente de conferência", "Rejeitado"]

# (chave, cabeçalho, largura, comentário)
COLUNAS = [
    ("id", "ID", 12, "Identificador único. O importador gera automaticamente (ex.: LAI-SS-000001). "
                     "Para entrada manual use CRI-<cidade>-<nº matrícula>-<ato>."),
    ("cidade", "Cidade", 15, "Escolha na lista."),
    ("bairro", "Bairro", 26, "Lista oficial da cidade escolhida (aba Bairros)."),
    ("endereco", "Logradouro e nº", 32, "Como consta no documento."),
    ("complemento", "Condomínio / Complemento", 24, "Unidade, bloco, condomínio, lote/quadra."),
    ("inscricao", "Inscrição imobiliária", 18, "Inscrição do IPTU, se constar."),
    ("tipo", "Tipo", 18, "Escolha na lista."),
    ("area_construida", "Área construída/privativa (m²)", 15, "Somente número. Vazio se o documento não informar."),
    ("area_terreno", "Área do terreno (m²)", 14, "Somente número. Vazio se o documento não informar."),
    ("padrao", "Padrão", 13, "Baixo / Normal / Alto, ou 'Não informado'. Critérios na aba Padrão."),
    ("origem_padrao", "Origem do padrão", 24, "De onde saiu a classificação do padrão."),
    ("padrao_original", "Padrão (texto do documento)", 22, "Texto literal do documento (ex.: classe do cadastro do IPTU)."),
    ("valor_venda", "Valor de venda (R$)", 16, "Valor declarado na escritura/registro (ato R de compra e venda) "
                                              "ou 'valor da transação' informado no ITBI. NUNCA o valor anunciado."),
    ("base_itbi", "Base de cálculo ITBI (R$)", 16, "Base de cálculo usada pela prefeitura, se informada."),
    ("venal_iptu", "Valor venal IPTU (R$)", 16, "Valor venal do IPTU no ano da transação, se informado."),
    ("rs_m2", "R$/m² (calculado)", 13, "Fórmula: valor de venda ÷ área construída/privativa "
                                      "(ou ÷ área do terreno, quando o tipo for Terreno). Não editar."),
    ("alerta", "Alerta", 22, "Fórmula. Sinaliza valor de venda menor ou igual ao venal do IPTU "
                            "(indício de subdeclaração). Não altera o valor."),
    ("data_escritura", "Data da escritura / transação", 14, "Data (dd/mm/aaaa)."),
    ("data_registro", "Data do registro", 14, "Data em que o ato foi registrado no CRI."),
    ("natureza", "Natureza", 16, "Somente compra e venda onerosa entra na base."),
    ("fonte", "Fonte", 24, "LAI ou certidão de matrícula."),
    ("documento", "Documento comprobatório", 30, "Matrícula nº + ato (ex.: M. 12.345 R-7) "
                                                 "ou protocolo LAI + linha do arquivo."),
    ("orgao", "CRI / Órgão", 26, "Cartório ou prefeitura que emitiu o documento."),
    ("coordenadas", "Coordenadas (lat,long)", 22, "Cole no formato -23.761234,-45.409876 "
                                                 "(clique direito no Google Maps copia neste formato)."),
    ("foto", "Foto da fachada", 22, "Preenchido pelo script vincular_fotos.py a partir da pasta fotos/."),
    ("origem_foto", "Origem da foto", 22, "Escolha na lista."),
    ("data_foto", "Data da foto", 12, "Lida do EXIF da foto quando existir."),
    ("street_view", "Street View", 16, "Fórmula: link do Street View a partir das coordenadas."),
    ("status", "Status", 20, "Validado = documento conferido. Só linhas validadas entram no Resumo."),
    ("observacoes", "Observações", 36, "Livre."),
]
CHAVES = [c[0] for c in COLUNAS]


def col(chave: str) -> str:
    """Letra da coluna da aba Vendas para uma chave."""
    return get_column_letter(CHAVES.index(chave) + 1)


def idx(chave: str) -> int:
    return CHAVES.index(chave) + 1


def normalizar(texto) -> str:
    """Minúsculas, sem acento e sem espaços duplicados (para comparar nomes)."""
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(s.lower().replace("-", " ").split())


# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
AZUL_ESCURO = "1F3864"
CINZA = "F2F2F2"
AMARELO = "FFF2CC"
fonte_normal = Font(name=FONTE, size=10)
fonte_cab = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
fonte_titulo = Font(name=FONTE, size=14, bold=True, color=AZUL_ESCURO)
fonte_negrito = Font(name=FONTE, size=10, bold=True)
fonte_link = Font(name=FONTE, size=10, color="0563C1", underline="single")
preench_cab = PatternFill("solid", fgColor=AZUL_ESCURO)
preench_formula = PatternFill("solid", fgColor=CINZA)
preench_entrada = PatternFill("solid", fgColor=AMARELO)
borda = Border(bottom=Side(style="thin", color="BFBFBF"))
quebra = Alignment(wrap_text=True, vertical="top")
FMT_REAIS = 'R$ #,##0.00;-R$ #,##0.00;"-"'
FMT_M2 = '#,##0.00'
FMT_DATA = 'DD/MM/YYYY'

COLUNAS_FORMULA = {"rs_m2", "alerta", "street_view"}


def _cabecalho(ws, titulos, linha=1, larguras=None):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font, c.fill = fonte_cab, preench_cab
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if larguras:
            ws.column_dimensions[get_column_letter(i)].width = larguras[i - 1]


def _texto(ws, linha, valor, negrito=False, coluna=1):
    c = ws.cell(row=linha, column=coluna, value=valor)
    c.font = fonte_negrito if negrito else fonte_normal
    c.alignment = quebra
    return c


# ---------------------------------------------------------------------------
# Fórmulas por linha
# ---------------------------------------------------------------------------
def formula_rs_m2(r: int) -> str:
    v, a, t, tipo = col("valor_venda"), col("area_construida"), col("area_terreno"), col("tipo")
    return (f'=IF(NOT(ISNUMBER({v}{r})),"",'
            f'IF({tipo}{r}="Terreno",IF(AND(ISNUMBER({t}{r}),{t}{r}>0),{v}{r}/{t}{r},""),'
            f'IF(AND(ISNUMBER({a}{r}),{a}{r}>0),{v}{r}/{a}{r},"")))')


def formula_alerta(r: int) -> str:
    v, n = col("valor_venda"), col("venal_iptu")
    return (f'=IF(AND(ISNUMBER({v}{r}),ISNUMBER({n}{r})),'
            f'IF({v}{r}<={n}{r},"Valor ≤ venal IPTU",""),"")')


def formula_street_view(r: int) -> str:
    c = col("coordenadas")
    return (f'=IF(LEN(TRIM({c}{r}))=0,"",HYPERLINK("https://www.google.com/maps/@?api=1&map_action=pano'
            f'&viewpoint="&SUBSTITUTE({c}{r}," ",""),"Abrir Street View"))')


def escrever_formulas_linha(ws, r: int):
    for chave, f in (("rs_m2", formula_rs_m2), ("alerta", formula_alerta), ("street_view", formula_street_view)):
        c = ws.cell(row=r, column=idx(chave), value=f(r))
        c.fill = preench_formula
        c.font = fonte_link if chave == "street_view" else fonte_normal


# ---------------------------------------------------------------------------
# Construção da planilha
# ---------------------------------------------------------------------------
def criar_planilha() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    _aba_vendas(ws)
    _aba_bairros(wb.create_sheet("Bairros"))
    _aba_listas(wb.create_sheet("Listas"))
    _aba_resumo(wb.create_sheet("Resumo"))
    _aba_padrao(wb.create_sheet("Padrão"))
    _aba_fontes(wb.create_sheet("Fontes"))
    _aba_regras(wb.create_sheet("Regras"))
    _aba_referencias(wb.create_sheet("Referências"))
    _aba_instrucoes(wb.create_sheet("Instruções", 0))
    wb.create_sheet("Galeria")
    reconstruir_galeria(wb, None)
    wb.calculation.fullCalcOnLoad = True
    return wb


def _aba_vendas(ws):
    _cabecalho(ws, [c[1] for c in COLUNAS], larguras=[c[2] for c in COLUNAS])
    for i, (_, _, _, nota) in enumerate(COLUNAS, start=1):
        ws.cell(row=1, column=i).comment = Comment(nota, "Base de vendas")
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ULTIMA_LINHA}"

    formatos = {"valor_venda": FMT_REAIS, "base_itbi": FMT_REAIS, "venal_iptu": FMT_REAIS,
                "rs_m2": FMT_REAIS, "area_construida": FMT_M2, "area_terreno": FMT_M2,
                "data_escritura": FMT_DATA, "data_registro": FMT_DATA, "data_foto": FMT_DATA}
    for r in range(PRIMEIRA_LINHA, ULTIMA_LINHA + 1):
        for chave, fmt in formatos.items():
            ws.cell(row=r, column=idx(chave)).number_format = fmt
        escrever_formulas_linha(ws, r)

    faixa = lambda chave: f"{col(chave)}{PRIMEIRA_LINHA}:{col(chave)}{ULTIMA_LINHA}"

    def lista(chave, nome_faixa, msg):
        dv = DataValidation(type="list", formula1=f"={nome_faixa}", allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor fora da lista", error=msg)
        dv.add(faixa(chave))
        ws.add_data_validation(dv)

    lista("cidade", "LISTA_CIDADES", "Escolha uma cidade da lista.")
    lista("tipo", "LISTA_TIPOS", "Escolha um tipo da lista.")
    lista("padrao", "LISTA_PADROES", "Use Baixo, Normal, Alto ou Não informado.")
    lista("origem_padrao", "LISTA_ORIGEM_PADRAO", "Escolha a origem do padrão.")
    lista("natureza", "LISTA_NATUREZAS", "Somente compra e venda onerosa entra na base.")
    lista("fonte", "LISTA_FONTES", "Somente LAI ou certidão de matrícula.")
    lista("origem_foto", "LISTA_ORIGEM_FOTO", "Escolha a origem da foto.")
    lista("status", "LISTA_STATUS", "Escolha um status.")

    c0 = col("cidade")
    dv_b = DataValidation(
        type="list", allow_blank=True, showErrorMessage=True, errorTitle="Bairro fora da lista oficial",
        error="Escolha primeiro a cidade; o bairro deve estar na lista oficial (aba Bairros).",
        formula1=f'=INDIRECT(IF(${c0}{PRIMEIRA_LINHA}="Bertioga","BAIRROS_BERTIOGA","BAIRROS_SAO_SEBASTIAO"))')
    dv_b.add(faixa("bairro"))
    ws.add_data_validation(dv_b)

    for chave in ("valor_venda", "base_itbi", "venal_iptu", "area_construida", "area_terreno"):
        dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Número inválido",
                            error="Informe apenas um número maior que zero (sem R$ ou m²).")
        dv.add(faixa(chave))
        ws.add_data_validation(dv)
    for chave in ("data_escritura", "data_registro", "data_foto"):
        dv = DataValidation(type="date", operator="between", formula1="DATE(1950,1,1)",
                            formula2="DATE(2100,12,31)", allow_blank=True, showErrorMessage=True,
                            errorTitle="Data inválida", error="Informe uma data válida (dd/mm/aaaa).")
        dv.add(faixa(chave))
        ws.add_data_validation(dv)

    alerta = f"{col('alerta')}{PRIMEIRA_LINHA}:{col('alerta')}{ULTIMA_LINHA}"
    ws.conditional_formatting.add(alerta, FormulaRule(
        formula=[f'LEN({col("alerta")}{PRIMEIRA_LINHA})>0'],
        fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=FONTE, color="9C0006")))
    st = f"{col('status')}{PRIMEIRA_LINHA}:{col('status')}{ULTIMA_LINHA}"
    for texto, cor in (("Validado", "C6EFCE"), ("Pendente de conferência", "FFEB9C"), ("Rejeitado", "FFC7CE")):
        ws.conditional_formatting.add(st, FormulaRule(
            formula=[f'{col("status")}{PRIMEIRA_LINHA}="{texto}"'], fill=PatternFill("solid", fgColor=cor)))


def _definir_nome(wb, nome, ref):
    wb.defined_names[nome] = DefinedName(nome, attr_text=ref)


def _aba_bairros(ws):
    ws["A1"] = "Bairros oficiais por cidade"
    ws["A1"].font = fonte_titulo
    linha = 3
    wb = ws.parent
    for coluna, cidade in ((1, "São Sebastião"), (5, "Bertioga")):
        letra = get_column_letter(coluna)
        _texto(ws, linha, cidade, True, coluna)
        _texto(ws, linha + 1, "Fonte: " + FONTE_BAIRROS[cidade], False, coluna)
        ws.merge_cells(start_row=linha + 1, start_column=coluna, end_row=linha + 1, end_column=coluna + 2)
        ws.row_dimensions[linha + 1].height = 60
        for j, t in enumerate(["Bairro", "Código", "Observação"]):
            c = ws.cell(row=linha + 2, column=coluna + j, value=t)
            c.font, c.fill = fonte_cab, preench_cab
        ini = linha + 3
        for k, (nome, cod, obs) in enumerate(BAIRROS[cidade]):
            for j, v in enumerate((nome, cod, obs)):
                ws.cell(row=ini + k, column=coluna + j, value=v).font = fonte_normal
        fim = ini + len(BAIRROS[cidade]) - 1
        _definir_nome(wb, NOME_FAIXA_BAIRROS[cidade], f"Bairros!${letra}${ini}:${letra}${fim}")
    for letra, w in zip("ABCDEFG", (28, 24, 48, 3, 36, 8, 48)):
        ws.column_dimensions[letra].width = w


def _aba_listas(ws):
    listas = [("Cidades", "LISTA_CIDADES", CIDADES), ("Tipos", "LISTA_TIPOS", TIPOS),
              ("Padrões", "LISTA_PADROES", PADROES), ("Origem do padrão", "LISTA_ORIGEM_PADRAO", ORIGEM_PADRAO),
              ("Natureza", "LISTA_NATUREZAS", NATUREZAS), ("Fontes", "LISTA_FONTES", FONTES),
              ("Origem da foto", "LISTA_ORIGEM_FOTO", ORIGEM_FOTO), ("Status", "LISTA_STATUS", STATUS)]
    for i, (titulo, nome, valores) in enumerate(listas, start=1):
        letra = get_column_letter(i)
        c = ws.cell(row=1, column=i, value=titulo)
        c.font, c.fill = fonte_cab, preench_cab
        for k, v in enumerate(valores, start=2):
            ws.cell(row=k, column=i, value=v).font = fonte_normal
        _definir_nome(ws.parent, nome, f"Listas!${letra}$2:${letra}${len(valores) + 1}")
        ws.column_dimensions[letra].width = 30


def _aba_resumo(ws):
    ws["A1"] = "Resumo por cidade e bairro (somente vendas com Status = Validado)"
    ws["A1"].font = fonte_titulo
    ws["A2"] = ("Fórmulas recalculam sozinhas. Bairros sem venda validada mostram '-'. "
                "Mediana e R$/m² ignoram linhas sem metragem.")
    ws["A2"].font = Font(name=FONTE, size=9, italic=True)
    titulos = ["Cidade", "Bairro", "Nº de vendas", "Valor mediano (R$)", "Menor valor (R$)",
               "Maior valor (R$)", "R$/m² mediano", "Nº c/ alerta venal", "Venda mais antiga", "Venda mais recente"]
    _cabecalho(ws, titulos, 4, [16, 30, 11, 18, 16, 16, 16, 12, 14, 14])

    R = lambda chave: f"Vendas!${col(chave)}${PRIMEIRA_LINHA}:${col(chave)}${ULTIMA_LINHA}"
    cid, bai, val, m2, st, al, dt = (R("cidade"), R("bairro"), R("valor_venda"), R("rs_m2"),
                                    R("status"), R("alerta"), R("data_escritura"))
    linha = 5
    linhas_cidade = {}
    for cidade in CIDADES:
        ini = linha
        for nome, _, _ in BAIRROS[cidade]:
            r = linha
            ws.cell(row=r, column=1, value=cidade)
            ws.cell(row=r, column=2, value=nome)
            cond = f'{cid},$A{r},{bai},$B{r},{st},"Validado"'
            arr = f'({cid}=$A{r})*({bai}=$B{r})*({st}="Validado")'
            ws.cell(row=r, column=3, value=f"=COUNTIFS({cond})")
            ws[f"D{r}"] = ArrayFormula(f"D{r}", f'=IF(C{r}=0,"-",MEDIAN(IF({arr}*ISNUMBER({val}),{val})))')
            ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"-",_xlfn.MINIFS({val},{cond}))')
            ws.cell(row=r, column=6, value=f'=IF(C{r}=0,"-",_xlfn.MAXIFS({val},{cond}))')
            ws[f"G{r}"] = ArrayFormula(
                f"G{r}", f'=IF(SUMPRODUCT({arr}*ISNUMBER({m2}))=0,"-",MEDIAN(IF({arr}*ISNUMBER({m2}),{m2})))')
            ws.cell(row=r, column=8, value=f'=COUNTIFS({cond},{al},"Valor*")')
            ws.cell(row=r, column=9, value=f'=IF(C{r}=0,"-",_xlfn.MINIFS({dt},{cond}))')
            ws.cell(row=r, column=10, value=f'=IF(C{r}=0,"-",_xlfn.MAXIFS({dt},{cond}))')
            linha += 1
        linhas_cidade[cidade] = (ini, linha - 1)
        r = linha
        ws.cell(row=r, column=1, value=cidade)
        ws.cell(row=r, column=2, value="TOTAL DA CIDADE")
        cond = f'{cid},$A{r},{st},"Validado"'
        arr = f'({cid}=$A{r})*({st}="Validado")'
        ws.cell(row=r, column=3, value=f"=COUNTIFS({cond})")
        ws[f"D{r}"] = ArrayFormula(f"D{r}", f'=IF(C{r}=0,"-",MEDIAN(IF({arr}*ISNUMBER({val}),{val})))')
        ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"-",_xlfn.MINIFS({val},{cond}))')
        ws.cell(row=r, column=6, value=f'=IF(C{r}=0,"-",_xlfn.MAXIFS({val},{cond}))')
        ws[f"G{r}"] = ArrayFormula(
            f"G{r}", f'=IF(SUMPRODUCT({arr}*ISNUMBER({m2}))=0,"-",MEDIAN(IF({arr}*ISNUMBER({m2}),{m2})))')
        ws.cell(row=r, column=8, value=f'=COUNTIFS({cond},{al},"Valor*")')
        ws.cell(row=r, column=9, value=f'=IF(C{r}=0,"-",_xlfn.MINIFS({dt},{cond}))')
        ws.cell(row=r, column=10, value=f'=IF(C{r}=0,"-",_xlfn.MAXIFS({dt},{cond}))')
        for cc in range(1, 11):
            ws.cell(row=r, column=cc).font = fonte_negrito
            ws.cell(row=r, column=cc).fill = preench_formula
        linha += 2

    # Por padrão construtivo
    linha += 1
    ws.cell(row=linha, column=1, value="Resumo por cidade e padrão").font = fonte_titulo
    linha += 1
    _cabecalho(ws, ["Cidade", "Padrão", "Nº de vendas", "Valor mediano (R$)", "Menor valor (R$)",
                    "Maior valor (R$)", "R$/m² mediano"], linha)
    linha += 1
    pad = R("padrao")
    for cidade in CIDADES:
        for p in PADROES:
            r = linha
            ws.cell(row=r, column=1, value=cidade)
            ws.cell(row=r, column=2, value=p)
            cond = f'{cid},$A{r},{pad},$B{r},{st},"Validado"'
            arr = f'({cid}=$A{r})*({pad}=$B{r})*({st}="Validado")'
            ws.cell(row=r, column=3, value=f"=COUNTIFS({cond})")
            ws[f"D{r}"] = ArrayFormula(f"D{r}", f'=IF(C{r}=0,"-",MEDIAN(IF({arr}*ISNUMBER({val}),{val})))')
            ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"-",_xlfn.MINIFS({val},{cond}))')
            ws.cell(row=r, column=6, value=f'=IF(C{r}=0,"-",_xlfn.MAXIFS({val},{cond}))')
            ws[f"G{r}"] = ArrayFormula(
                f"G{r}", f'=IF(SUMPRODUCT({arr}*ISNUMBER({m2}))=0,"-",MEDIAN(IF({arr}*ISNUMBER({m2}),{m2})))')
            linha += 1

    for row in ws.iter_rows(min_row=5, max_row=linha):
        for c in row:
            if c.font != fonte_negrito and not (c.font and c.font.bold):
                c.font = fonte_normal
            if c.column in (4, 5, 6, 7):
                c.number_format = FMT_REAIS
            elif c.column in (9, 10):
                c.number_format = FMT_DATA
    ws.freeze_panes = "C5"


def _aba_padrao(ws):
    ws["A1"] = "Critério de padrão construtivo"
    ws["A1"].font = fonte_titulo
    linhas = [
        ("Regra geral", "O padrão só é preenchido a partir de documento ou vistoria. Sem informação, "
                        "use 'Não informado'. Nunca deduzir pelo preço, bairro ou anúncio."),
        ("1ª opção — Cadastro municipal", "Se o ITBI/IPTU trouxer a classe/tipo/padrão construtivo da prefeitura, "
                                          "copie o texto literal em 'Padrão (texto do documento)' e converta para "
                                          "Baixo/Normal/Alto com a tabela de equivalência do importador "
                                          "(arquivo mapa_colunas.json, chave 'mapa_padrao'). Origem: Cadastro municipal."),
        ("2ª opção — Documento", "Memorial descritivo, habite-se ou incorporação registrada na matrícula que "
                                 "declare o padrão (a NBR 12721 exige que incorporações informem o padrão: "
                                 "baixo, normal ou alto). Origem: Certidão/memorial/habite-se."),
        ("3ª opção — Vistoria própria", "Classificação feita por vistoria do imóvel seguindo os projetos-padrão "
                                        "da ABNT NBR 12721 (Baixo / Normal / Alto), registrando a data e quem "
                                        "vistoriou em Observações. Origem: Vistoria própria."),
        ("Baixo", "NBR 12721 — padrão de acabamento baixo."),
        ("Normal", "NBR 12721 — padrão de acabamento normal."),
        ("Alto", "NBR 12721 — padrão de acabamento alto."),
        ("Não informado", "O documento não informa e não houve vistoria."),
    ]
    _cabecalho(ws, ["Item", "Critério"], 3, [30, 110])
    for i, (a, b) in enumerate(linhas, start=4):
        _texto(ws, i, a, True, 1)
        _texto(ws, i, b, False, 2)


def _aba_fontes(ws):
    ws["A1"] = "Registro dos documentos recebidos"
    ws["A1"].font = fonte_titulo
    titulos = ["Código do documento", "Tipo", "Órgão / Cartório", "Nº protocolo / matrícula",
               "Data do pedido", "Data da resposta/emissão", "Arquivo (caminho)", "Custo (R$)", "Observações"]
    _cabecalho(ws, titulos, 3, [20, 22, 30, 24, 14, 16, 36, 12, 40])
    dv = DataValidation(type="list", formula1='"Resposta LAI,Recurso LAI,Certidão de matrícula,Outro"',
                        allow_blank=True)
    dv.add("B4:B1000")
    ws.add_data_validation(dv)
    for r in range(4, 1001):
        ws.cell(row=r, column=5).number_format = FMT_DATA
        ws.cell(row=r, column=6).number_format = FMT_DATA
        ws.cell(row=r, column=8).number_format = FMT_REAIS
    ws.freeze_panes = "A4"


REGRAS = [
    ("1. Valor real", "Valor de venda = valor declarado na escritura registrada (ato R de compra e venda na "
                      "matrícula) ou valor da transação informado pela prefeitura no ITBI. Valor anunciado, "
                      "avaliação de corretor ou estimativa NUNCA entram."),
    ("2. Natureza", "Só compra e venda onerosa. Doação, herança/inventário, permuta, integralização de capital, "
                    "cessão, dação em pagamento, adjudicação, arrematação/leilão e consolidação de propriedade "
                    "fiduciária ficam fora (o importador lista como rejeitadas)."),
    ("3. Valor e base de cálculo", "Valor de venda e base de cálculo do ITBI ficam em colunas separadas. "
                                   "Um não substitui o outro."),
    ("4. Subdeclaração", "Quando o valor de venda é menor ou igual ao valor venal do IPTU a coluna Alerta "
                         "sinaliza. O valor NÃO é corrigido: o alerta só indica que o valor declarado pode "
                         "estar abaixo do praticado."),
    ("5. Duplicidade", "A mesma venda (mesma cidade + matrícula ou inscrição + data) não pode aparecer duas "
                       "vezes, mesmo que venha da LAI e da certidão. Mantém-se a linha da certidão e cita-se o "
                       "protocolo LAI em Observações."),
    ("6. Metragem", "Sem área no documento, a área fica vazia e o R$/m² também. Nunca estimar."),
    ("7. Bairro", "Só bairros da lista oficial (aba Bairros). Nomes populares são convertidos pelo "
                  "importador via 'apelidos_bairro' do mapa_colunas.json, sempre de forma explícita."),
    ("8. Padrão", "Conforme aba Padrão. Sem documento ou vistoria: 'Não informado'."),
    ("9. Foto", "Somente foto real do imóvel: tirada no local, link do Street View nas coordenadas do "
                "imóvel ou imagem constante de documento. Nunca imagem de anúncio de outro imóvel."),
    ("10. Status", "Validado = documento conferido e arquivado (aba Fontes). Pendente = falta conferir. "
                   "Rejeitado = não atende às regras. O Resumo usa só Validado."),
]


def _aba_regras(ws):
    ws["A1"] = "Regras de validação"
    ws["A1"].font = fonte_titulo
    _cabecalho(ws, ["Regra", "Descrição"], 3, [26, 120])
    for i, (a, b) in enumerate(REGRAS, start=4):
        _texto(ws, i, a, True, 1)
        _texto(ws, i, b, False, 2)


REFERENCIAS = [
    ("Bairros de Bertioga", "LC 99/2013 de Bertioga, Anexo 1",
     "https://www.bertioga.sp.gov.br/wp/wp-content/uploads/2015/06/Lei-Complementar-99.2013.pdf"),
    ("Bairros de São Sebastião", "IBGE – malha de bairros do Censo 2022 (SP)",
     "https://geoftp.ibge.gov.br/organizacao_do_territorio/malhas_territoriais/malhas_de_setores_censitarios__divisoes_intramunicipais/censo_2022/bairros/shp/UF/SP_bairros_CD2022.zip"),
    ("CRI São Sebastião", "Oficial de Registro de Imóveis e Anexos de São Sebastião",
     "https://www.risaosebastiao.com.br"),
    ("CRI Bertioga (desde 01/04/2025)", "Oficial de Registro de Imóveis de Bertioga — criado pela Lei Estadual "
     "nº 18.075/2024", "https://registrobertioga.com.br/"),
    ("Lei de criação do CRI Bertioga", "Lei Estadual nº 18.075, de 27/12/2024",
     "https://www.al.sp.gov.br/repositorio/legislacao/lei/2024/lei-18075-27.12.2024.html"),
    ("Matrículas antigas de Bertioga", "1º Oficial de Registro de Imóveis de Santos (atendia Bertioga até a "
     "instalação do CRI local)", "https://www.1risantos.com.br/"),
    ("Certidão online", "RI Digital / SAEC (ONR) — antigo registradores.onr.org.br", "https://ridigital.org.br/"),
    ("e-SIC São Sebastião", "Sistema de Informação ao Cidadão da Prefeitura",
     "http://www.saosebastiao.sp.gov.br/esic/"),
    ("Transparência Bertioga", "Portal da Transparência (seção e-SIC)",
     "https://transparencia-bertioga.smarapd.com.br/#/esic"),
    ("Ouvidoria Bertioga", "Fala Cidadão", "https://sistemas-smarapd.bertioga.sp.gov.br/falacidadao/#!/demanda"),
    ("Código Tributário de Bertioga", "LC 116/2015 (art. 82 – base de cálculo do ITBI)",
     "https://sapl.bertioga.sp.leg.br/media/sapl/public/normajuridica/2015/3823/3823_texto_integral.pdf"),
    ("Precedente – capital SP", "A Prefeitura de São Paulo publica as transações de ITBI sem nomes/CPF",
     "https://www.prefeitura.sp.gov.br/cidade/secretarias/fazenda/noticias/index.php?p=31551"),
    ("Lei de Acesso à Informação", "Lei nº 12.527/2011", "https://www.planalto.gov.br/ccivil_03/_ato2011-2014/2011/lei/l12527.htm"),
]


def _aba_referencias(ws):
    ws["A1"] = "Referências e canais"
    ws["A1"].font = fonte_titulo
    _cabecalho(ws, ["Assunto", "Descrição", "Link"], 3, [30, 70, 90])
    for i, (a, b, url) in enumerate(REFERENCIAS, start=4):
        _texto(ws, i, a, True, 1)
        _texto(ws, i, b, False, 2)
        c = ws.cell(row=i, column=3, value=url)
        c.hyperlink = url
        c.font = fonte_link
    _texto(ws, len(REFERENCIAS) + 5,
           f"Links verificados em {datetime.now():%d/%m/%Y}. O site da Prefeitura de São Sebastião estava "
           "fora do ar durante a verificação; confirme o endereço do e-SIC antes de protocolar.")


def _aba_instrucoes(ws):
    ws["A1"] = "Base de imóveis vendidos — São Sebastião e Bertioga"
    ws["A1"].font = fonte_titulo
    textos = [
        ("O que é", "Base de vendas efetivas de imóveis com o valor REAL de venda (escritura/registro ou ITBI), "
                    "nunca o valor anunciado. Cada linha precisa de um documento comprobatório."),
        ("Por que começa vazia", "Não existem dados públicos de vendas para estas cidades. Os dados entram à medida "
                                 "que chegam as respostas dos pedidos LAI e as certidões de matrícula."),
        ("Onde digitar", "Aba Vendas. Células comuns = entrada. Células cinza = fórmulas (não editar): R$/m², "
                         "Alerta e Street View. Passe o mouse no cabeçalho para ver a explicação de cada coluna."),
        ("Importar LAI", "python scripts/importar_lai.py <arquivo_da_prefeitura> <mapa_colunas.json>"),
        ("Fotos", "Salve as fotos em fotos/ com o nome do ID (ex.: fotos/CRI-SS-12345-R7.jpg; várias: "
                  "..._2.jpg) e rode python scripts/vincular_fotos.py. A coluna 'Foto da fachada' recebe o link "
                  "e a aba Galeria mostra as miniaturas."),
        ("Resumo", "Aba Resumo: por cidade/bairro e por padrão, só com linhas Validado."),
        ("Regras", "Aba Regras. Critério de padrão na aba Padrão. Documentos recebidos na aba Fontes."),
    ]
    _cabecalho(ws, ["Tópico", "Descrição"], 3, [24, 120])
    r = 4
    for a, b in textos:
        _texto(ws, r, a, True, 1)
        _texto(ws, r, b, False, 2)
        r += 1
    r += 1
    _texto(ws, r, "Exemplo de preenchimento (FICTÍCIO — só mostra o formato; não é dado real e não está na aba "
                  "Vendas)", True)
    r += 1
    exemplo = {
        "id": "CRI-SS-00000-R0", "cidade": "São Sebastião", "bairro": "Maresias",
        "endereco": "Rua Exemplo, 000", "complemento": "Cond. Exemplo, casa 0", "inscricao": "0000.000.000",
        "tipo": "Casa em condomínio", "area_construida": "000,00", "area_terreno": "000,00",
        "padrao": "Alto", "origem_padrao": "Certidão/memorial/habite-se", "padrao_original": "texto literal",
        "valor_venda": "0.000.000,00", "base_itbi": "0.000.000,00", "venal_iptu": "000.000,00",
        "rs_m2": "(automático)", "alerta": "(automático)", "data_escritura": "dd/mm/aaaa",
        "data_registro": "dd/mm/aaaa", "natureza": "Compra e venda", "fonte": "Certidão de matrícula (CRI)",
        "documento": "Matrícula 00.000 — R-0", "orgao": "RI São Sebastião",
        "coordenadas": "-23.000000,-45.000000", "foto": "(automático)", "origem_foto": "Foto própria no local",
        "data_foto": "dd/mm/aaaa", "street_view": "(automático)", "status": "Validado", "observacoes": "",
    }
    for i, (chave, titulo, _, _) in enumerate(COLUNAS):
        _texto(ws, r + i, titulo, True, 1)
        c = _texto(ws, r + i, exemplo[chave], False, 2)
        c.font = Font(name=FONTE, size=10, italic=True, color="7F7F7F")


# ---------------------------------------------------------------------------
# Galeria de fotos
# ---------------------------------------------------------------------------
EXT_FOTO = (".jpg", ".jpeg", ".png")


def fotos_do_id(pasta_base: str, id_: str) -> list[str]:
    """Arquivos em fotos/ com nome <ID>.ext ou <ID>_<n>.ext (caminhos relativos)."""
    pasta = os.path.join(pasta_base, PASTA_FOTOS)
    if not id_ or not os.path.isdir(pasta):
        return []
    achados = []
    for nome in sorted(os.listdir(pasta)):
        raiz, ext = os.path.splitext(nome)
        if ext.lower() in EXT_FOTO and (raiz == id_ or raiz.startswith(id_ + "_")):
            achados.append(os.path.join(PASTA_FOTOS, nome))
    return achados


def reconstruir_galeria(wb, pasta_base: str | None):
    """Recria a aba Galeria com miniaturas das fotos vinculadas.

    O openpyxl não preserva imagens ao reabrir um arquivo, por isso a galeria é
    sempre refeita a partir da pasta fotos/ antes de salvar.
    """
    if "Galeria" in wb.sheetnames:
        pos = wb.sheetnames.index("Galeria")
        del wb["Galeria"]
        ws = wb.create_sheet("Galeria", pos)
    else:
        ws = wb.create_sheet("Galeria")
    ws["A1"] = "Galeria de fachadas"
    ws["A1"].font = fonte_titulo
    ws["A2"] = "Gerada por scripts/vincular_fotos.py a partir da pasta fotos/. Não editar à mão."
    ws["A2"].font = Font(name=FONTE, size=9, italic=True)
    _cabecalho(ws, ["ID", "Cidade", "Bairro", "Endereço", "Arquivo", "Miniatura"], 4, [18, 15, 22, 32, 30, 34])
    if pasta_base is None:
        return 0
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    vendas = wb["Vendas"]
    pasta_mini = os.path.join(pasta_base, PASTA_FOTOS, ".miniaturas")
    linha, total = 5, 0
    for r in range(PRIMEIRA_LINHA, vendas.max_row + 1):
        id_ = vendas.cell(row=r, column=idx("id")).value
        if not id_:
            continue
        for rel in fotos_do_id(pasta_base, str(id_)):
            os.makedirs(pasta_mini, exist_ok=True)
            mini = os.path.join(pasta_mini, os.path.basename(rel).rsplit(".", 1)[0] + ".png")
            with PILImage.open(os.path.join(pasta_base, rel)) as im:
                im = _orientar(im)
                im.thumbnail((240, 180))
                im.convert("RGB").save(mini, "PNG")
            valores = [id_, vendas.cell(row=r, column=idx("cidade")).value,
                       vendas.cell(row=r, column=idx("bairro")).value,
                       vendas.cell(row=r, column=idx("endereco")).value, rel]
            for j, v in enumerate(valores, start=1):
                c = ws.cell(row=linha, column=j, value=v)
                c.font, c.alignment = fonte_normal, quebra
            ws.cell(row=linha, column=5).hyperlink = rel.replace(os.sep, "/")
            ws.cell(row=linha, column=5).font = fonte_link
            img = XLImage(mini)
            ws.add_image(img, f"F{linha}")
            ws.row_dimensions[linha].height = 140
            linha += 1
            total += 1
    return total


def _orientar(im):
    try:
        from PIL import ImageOps
        return ImageOps.exif_transpose(im)
    except Exception:
        return im


def salvar(wb, caminho: str):
    """Reconstrói a galeria (imagens se perdem ao reabrir) e salva."""
    reconstruir_galeria(wb, os.path.dirname(os.path.abspath(caminho)))
    wb.calculation.fullCalcOnLoad = True
    wb.save(caminho)
